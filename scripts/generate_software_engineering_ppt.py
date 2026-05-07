# -*- coding: utf-8 -*-
"""Generate a Software Engineering overview PowerPoint with a cohesive theme."""
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

# --- Theme: 「深蓝工程」Engineering Blueprint ---
# 主色：深海蓝 | 强调：青蓝 | 背景：云灰 | 正文：板岩灰
THEME = {
    "primary": RGBColor(0x0F, 0x17, 0x2A),      # slate-900 标题栏 / 封面底
    "primary_light": RGBColor(0x1E, 0x3A, 0x5F),  # 略浅，装饰用
    "accent": RGBColor(0x22, 0xD3, 0xEE),       # cyan-400 点缀
    "accent_muted": RGBColor(0x67, 0xE8, 0xF9), # 副标题 / 高亮
    "page": RGBColor(0xF8, 0xFA, 0xFC),         # slate-50 内容页底
    "text": RGBColor(0x33, 0x41, 0x55),         # slate-700 正文
    "text_dim": RGBColor(0x64, 0x74, 0x8B),     # slate-500 次要
    "white": RGBColor(0xFF, 0xFF, 0xFF),
}

FONT_CN = "微软雅黑"
TITLE_BAR_H = Inches(1.05)
BODY_TOP = Inches(1.18)


def _set_fill_solid(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    try:
        shape.line.fill.background()
    except Exception:
        pass


def style_opening_slide(slide, prs):
    """封面 / 封底：深色全屏 + 白字 + 青蓝强调。"""
    bg = slide.background
    bg.fill.solid()
    bg.fill.fore_color.rgb = THEME["primary"]

    title = slide.shapes.title
    title_text = title.text
    tf = title.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.name = FONT_CN
    p.font.size = Pt(40)
    p.font.bold = True
    p.font.color.rgb = THEME["white"]

    sub = slide.placeholders[1]
    stf = sub.text_frame
    for para in stf.paragraphs:
        para.font.name = FONT_CN
        para.font.size = Pt(18)
        para.font.color.rgb = THEME["accent_muted"]

    title.left = Inches(0.6)
    title.top = Inches(2.4)
    title.width = prs.slide_width - Inches(1.2)
    title.height = Inches(1.2)

    sub.left = Inches(0.6)
    sub.top = Inches(3.85)
    sub.width = prs.slide_width - Inches(1.2)
    sub.height = Inches(2.5)


def style_content_slide(slide, prs):
    """内容页：浅底 + 全宽深蓝标题栏 + 正文区。"""
    bg = slide.background
    bg.fill.solid()
    bg.fill.fore_color.rgb = THEME["page"]

    title = slide.shapes.title
    title.left = 0
    title.top = 0
    title.width = prs.slide_width
    title.height = TITLE_BAR_H
    _set_fill_solid(title, THEME["primary"])

    tf = title.text_frame
    tf.word_wrap = True
    for para in tf.paragraphs:
        para.font.name = FONT_CN
        para.font.size = Pt(26)
        para.font.bold = True
        para.font.color.rgb = THEME["white"]
        para.alignment = PP_ALIGN.CENTER

    tf.vertical_anchor = MSO_ANCHOR.MIDDLE

    body = slide.placeholders[1]
    body.left = Inches(0.55)
    body.top = BODY_TOP
    body.width = prs.slide_width - Inches(1.1)
    body.height = prs.slide_height - BODY_TOP - Inches(0.55)

    btf = body.text_frame
    btf.word_wrap = True
    for para in btf.paragraphs:
        para.font.name = FONT_CN
        para.font.size = Pt(19)
        para.font.color.rgb = THEME["text"]
        para.space_after = Pt(10)


def add_bullet_slide(prs, title, bullets):
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = title
    body = slide.placeholders[1]
    tf = body.text_frame
    tf.clear()
    for i, line in enumerate(bullets):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = line
        p.level = 0
    style_content_slide(slide, prs)
    return slide


def main():
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)

    # --- Title ---
    slide0 = prs.slides.add_slide(prs.slide_layouts[0])
    slide0.shapes.title.text = "软件工程技术"
    slide0.placeholders[1].text = "概念、流程与实践概览\n\n适合：课程复习 / 团队分享 / 入门导读"
    style_opening_slide(slide0, prs)

    # --- Outline ---
    add_bullet_slide(prs, "目录", [
        "软件工程与目标",
        "生命周期与过程模型",
        "需求与建模",
        "设计与架构",
        "实现、质量与测试",
        "运维与工程文化",
        "趋势与小结",
    ])

    # --- Section 1 ---
    add_bullet_slide(prs, "软件工程是什么？", [
        "将系统化、可度量的方法用于软件的开发、运行与维护",
        "目标：在约束（时间、成本、质量）下交付可维护、可演进的系统",
        "核心活动：需求 → 设计 → 实现 → 验证 → 运维 → 退役",
        "与「只写代码」的区别：强调过程、协作、风险与可重复性",
    ])

    add_bullet_slide(prs, "质量属性（非功能需求）", [
        "可靠性、可用性、性能、安全、可维护性、可扩展性",
        "可测试性、可移植性、易用性（用户体验）",
        "工程上需在需求阶段明确优先级，避免后期无法弥补的架构债",
    ])

    # --- Section 2 ---
    add_bullet_slide(prs, "生命周期与过程模型", [
        "瀑布模型：阶段清晰，适合需求稳定；变更成本高",
        "增量 / 迭代：分阶段交付，降低风险",
        "敏捷（Scrum、Kanban）：短周期、持续反馈、拥抱变化",
        "DevOps：开发运维一体化，强调自动化与快速反馈",
    ])

    add_bullet_slide(prs, "敏捷实践要点", [
        "用户故事与待办列表（Product Backlog）",
        "迭代（Sprint）计划、每日站会、评审与回顾",
        "持续集成：小步提交、自动构建与测试",
        "可工作软件优于详尽文档（在合适粒度上平衡文档）",
    ])

    # --- Section 3 ---
    add_bullet_slide(prs, "需求工程", [
        "获取：访谈、原型、观察、竞品分析",
        "分析：用例、用户故事、业务规则",
        "规格说明：清晰、可验证、可追踪（到设计与测试）",
        "变更管理：版本、评审、影响分析",
    ])

    add_bullet_slide(prs, "建模（简要）", [
        "结构化：数据流图、实体关系图",
        "面向对象：用例图、类图、时序图、状态图（UML）",
        "目的：沟通共识、发现遗漏、支撑设计与测试",
    ])

    # --- Section 4 ---
    add_bullet_slide(prs, "软件设计原则", [
        "模块化、高内聚、低耦合",
        "SOLID：单一职责、开闭、里氏替换、接口隔离、依赖倒置",
        "DRY、KISS、YAGNI：避免重复、保持简单、不过度设计",
        "设计模式：在合适场景复用成熟方案（非炫技）",
    ])

    add_bullet_slide(prs, "架构与常见风格", [
        "分层架构、客户端-服务器、微服务（与运维复杂度权衡）",
        "事件驱动、消息队列：解耦与异步扩展",
        "API 设计：REST/GraphQL 等，关注版本化与兼容性",
        "数据：一致性、事务边界、缓存与分库分表策略",
    ])

    # --- Section 5 ---
    add_bullet_slide(prs, "实现与代码质量", [
        "编码规范与 Code Review：风格一致、知识传递",
        "静态分析、类型系统、重构：控制技术债",
        "版本控制（Git）：分支策略、合并请求、可追溯性",
        "安全：输入校验、认证授权、依赖与密钥管理",
    ])

    add_bullet_slide(prs, "测试层次", [
        "单元测试：函数/类级别，快速反馈",
        "集成测试：模块协作、接口契约",
        "系统 / 端到端测试：用户路径与关键场景",
        "测试金字塔：底层多、顶层少；自动化与手工探索性测试结合",
    ])

    # --- Section 6 ---
    add_bullet_slide(prs, "运维与工程效能", [
        "持续集成 / 持续交付（CI/CD）：流水线、环境一致性",
        "可观测性：日志、指标、链路追踪",
        "发布策略：蓝绿、金丝雀、特性开关",
        "SRE 思想：错误预算、SLI/SLO、事后复盘（Blameless）",
    ])

    add_bullet_slide(prs, "团队协作与文档", [
        "角色：产品、开发、测试、运维、安全（DevSecOps）",
        "知识库：架构决策记录（ADR）、Runbook、API 文档",
        "沟通：站会、评审、跨职能协作",
    ])

    # --- Section 7 ---
    add_bullet_slide(prs, "技术趋势（选修）", [
        "云原生：容器、Kubernetes、Serverless",
        "平台工程：内部开发者平台（IDP）、自服务基础设施",
        "AI 辅助：代码补全、测试生成、需人工审核与规范约束",
        "低代码 / 领域驱动设计（DDD）在复杂业务中的应用",
    ])

    add_bullet_slide(prs, "小结", [
        "软件工程 = 技术 + 过程 + 人 + 工具",
        "没有银弹：按场景选择模型与实践，持续度量与改进",
        "质量左移：需求与设计阶段多投入，降低后期成本",
        "推荐延伸：SWEBOK、敏捷宣言、《人月神话》《架构整洁之道》等",
    ])

    # --- Thank you ---
    last = prs.slides.add_slide(prs.slide_layouts[0])
    last.shapes.title.text = "谢谢"
    last.placeholders[1].text = "问答与交流"
    style_opening_slide(last, prs)

    # 若保存失败（文件被 PowerPoint 占用），请先关闭 pptx 再运行，或临时改路径
    out = r"e:\mechantuni\mechant\software-engineering-tech.pptx"
    prs.save(out)
    print("Saved:", out)


if __name__ == "__main__":
    main()
